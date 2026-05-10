from docx import Document
import os

TEMPLATE_DIR = r"E:\vishnu viswas\ai_valuation\templates"
os.makedirs(TEMPLATE_DIR, exist_ok=True)

doc = Document()

# -----------------------------
# TITLE
# -----------------------------
doc.add_heading("VALUATION REPORT", level=1)

# -----------------------------
# BASIC DETAILS
# -----------------------------
doc.add_paragraph("Borrower Name : {{ BORROWER_NAME }}")
doc.add_paragraph("Survey No     : {{ SURVEY_NO }}")
doc.add_paragraph("Village       : {{ VILLAGE }}")
doc.add_paragraph("Taluk         : {{ TALUK }}")
doc.add_paragraph("District      : {{ DISTRICT }}")

# -----------------------------
# BOUNDARIES
# -----------------------------
doc.add_paragraph("")
doc.add_heading("Boundaries of the Property", level=2)

doc.add_paragraph("East  : {{ BOUNDARY_EAST }}")
doc.add_paragraph("North : {{ BOUNDARY_NORTH }}")
doc.add_paragraph("West  : {{ BOUNDARY_WEST }}")
doc.add_paragraph("South : {{ BOUNDARY_SOUTH }}")

# -----------------------------
# LAND DETAILS
# -----------------------------
doc.add_paragraph("")
doc.add_heading("Land Details", level=2)

doc.add_paragraph("Extent (Ares)         : {{ EXTENT }}")
doc.add_paragraph("Total Cents          : {{ TOTAL_CENTS }}")
doc.add_paragraph("Land Rate / Cent (₹) : {{ LAND_RATE_PER_CENT }}")
doc.add_paragraph("Raw Land Value (₹)   : {{ RAW_LAND_VALUE }}")
doc.add_paragraph("Total Land Value (₹) : {{ TOTAL_LAND_VALUE }}")

# -----------------------------
# BUILDING DETAILS
# -----------------------------
doc.add_paragraph("")
doc.add_heading("Building Details", level=2)

doc.add_paragraph("Building Area (Sqft)        : {{ BUILDING_AREA_SQFT }}")
doc.add_paragraph("Building Rate / Sqft (₹)    : {{ BUILDING_RATE_PER_SQFT }}")
doc.add_paragraph("Building Age (Years)        : {{ BUILDING_AGE_YEARS }}")
doc.add_paragraph("Gross Building Value (₹)    : {{ BUILDING_GROSS_VALUE }}")
doc.add_paragraph("Building Depreciation (₹)   : {{ BUILDING_DEPRECIATION }}")
doc.add_paragraph("Depreciated Building Value (₹) : {{ DEPRECIATED_BUILDING_VALUE }}")

# -----------------------------
# TOTAL VALUES
# -----------------------------
doc.add_paragraph("")
doc.add_heading("Valuation Summary", level=2)

doc.add_paragraph("Market Value (₹)     : {{ MARKET_VALUE }}")
doc.add_paragraph("Market Value (Words): {{ MARKET_VALUE_WORDS }}")

doc.add_paragraph("Realizable Value (₹)     : {{ REALIZABLE_VALUE }}")
doc.add_paragraph("Realizable Value (Words): {{ REALIZABLE_VALUE_WORDS }}")

doc.add_paragraph("Distress Value (₹)     : {{ DISTRESS_VALUE }}")
doc.add_paragraph("Distress Value (Words): {{ DISTRESS_VALUE_WORDS }}")

# -----------------------------
# META
# -----------------------------
doc.add_paragraph("")
doc.add_paragraph("Place : {{ PLACE }}")
doc.add_paragraph("Date  : {{ REPORT_DATE }}")

# -----------------------------
# SAVE TEMPLATE
# -----------------------------
file_path = os.path.join(TEMPLATE_DIR, "valuation_template.docx")
doc.save(file_path)

print("✅ Clean valuation template created at:")
print(file_path)
