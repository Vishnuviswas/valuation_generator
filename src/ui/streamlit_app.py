import streamlit as st
import os
import sys
import re
import json
import ollama
from num2words import num2words
from datetime import datetime,timedelta
from docx import Document

# -------------------------------------------------
# PATH FIX
# -------------------------------------------------
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "../.."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from src.pdf_utils import pdf_to_images
from src.ocr.tesseract_english import ocr_english
from src.nlp.field_extractor import extract_all_fields
from src.report.gen_report import generate_report

# -------------------------------------------------
# CONSTANTS
# -------------------------------------------------
ARE_TO_CENT = 2.471
DEPRECIATION_PER_YEAR = 0.02

# -------------------------------------------------
# HELPERS
# -------------------------------------------------
def to_words(amount):
    return num2words(int(amount), lang="en_IN").title() + " Only"

def round_down(value, unit):
    return int(value // unit) * unit

def read_docx_text(docx_path):
    doc = Document(docx_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

# -------------------------------------------------
# AI FIELD EXTRACTION (OLLAMA - LOCAL, OFFLINE)
# -------------------------------------------------
AI_MODEL = "llama3.2:3b"

AI_FIELDS = ["borrower", "owner", "survey_no", "village", "taluk", "district",
             "door_no", "boundary_east", "boundary_west",
             "boundary_north", "boundary_south", "extent",
             "sale_deed", "location_sketch", "tax_receipt", "possession_cert"]

AI_SCHEMA = {
    "type": "object",
    "properties": {k: {"type": "string"} for k in AI_FIELDS},
    "required": AI_FIELDS,
}

def ai_extract_fields(text):
    prompt = (
        "You are reading OCR text of a Kerala property title investigation report.\n"
        "Extract the fields exactly as written in the document.\n"
        "survey_no: include resurvey and block numbers if given.\n"
        "extent: copy with its unit (Ares / cents / hectares).\n"
        "sale_deed: sale deed number, date and SRO, e.g. '4305/2015 dated 27.11.2015 at Paravur SRO'.\n"
        "location_sketch: location sketch number and village, e.g. '1120/2026 of Paravur Village'.\n"
        "tax_receipt: land tax receipt number and village, e.g. 'KL07030304022/2026 of Paravur Village'.\n"
        "possession_cert: possession certificate number and village, e.g. '101469411 of Paravur Village'.\n"
        "If a field is not present, return an empty string. Do not guess.\n\n"
        "TEXT:\n" + text[:8000]
    )
    r = ollama.chat(
        model=AI_MODEL,
        format=AI_SCHEMA,
        keep_alive=0,
        options={"temperature": 0, "num_ctx": 4096},
        messages=[{"role": "user", "content": prompt}],
    )
    return json.loads(r["message"]["content"])

def extract_documents_regex(text):
    """Backup: find legal document numbers by their fixed keywords."""
    pats = {
        "sale_deed": r"(?:Sale\s*Deed|Document)\s*No\.?\s*[:\-]?\s*(\d+\s*/\s*\d{4}[^\n]{0,60})",
        "location_sketch": r"Location\s*Sketch\s*No\.?\s*[:\-]?\s*([\w/]+[^\n]{0,40})",
        "tax_receipt": r"(?:Land\s*)?Tax\s*Receipt\s*No\.?\s*[:\-]?\s*([A-Z0-9/]+[^\n]{0,40})",
        "possession_cert": r"Possession\s*Certificate\s*No\.?\s*[:\-]?\s*([\w/]+[^\n]{0,40})",
    }
    out = {}
    for k, p in pats.items():
        m = re.search(p, text, re.I)
        out[k] = m.group(1).strip(" .,") if m else ""
    return out

def extract_fields_with_ai(text):
    """AI first; rule-based extractor and regex fill any fields the AI left empty."""
    try:
        fields = ai_extract_fields(text)
    except Exception as e:
        st.warning(f"AI extraction failed, using rules only: {e}")
        fields = {}
    rule_fields = extract_all_fields(text)
    for k, v in rule_fields.items():
        if v and not fields.get(k):
            fields[k] = v
    for k, v in extract_documents_regex(text).items():
        if v and not fields.get(k):
            fields[k] = v
    return fields

# -------------------------------------------------
# FEDERAL BANK LOGIC
# -------------------------------------------------
def calculate_federal_separate(land_value, building_value):
    land_market = round_down(land_value, 100000)
    land_realizable = round_down(land_market * 0.90, 100000)
    land_distress = round_down(land_market * 0.80, 100000)



    building_market = round_down(building_value, 100000)
    building_realizable = round_down(building_market * 0.90, 100000)
    building_distress = round_down(building_market * 0.80, 100000)

    return {
        "LAND_MARKET_VALUE": format_indian_number(land_market),
        "LAND_REALIZABLE_VALUE": format_indian_number(land_realizable),
        "LAND_DISTRESS_VALUE": format_indian_number(land_distress),
        "BUILDING_MARKET_VALUE": format_indian_number(building_market),
        "BUILDING_GROSS_VALUE":gross_building_value,
        "BUILDING_REALIZABLE_VALUE":format_indian_number(building_realizable),
        "BUILDING_DISTRESS_VALUE":format_indian_number( building_distress),
        "MARKET_VALUE": land_market + building_market,
        "REALIZABLE_VALUE": land_realizable + building_realizable,
        "DISTRESS_VALUE": land_distress + building_distress,
    }
from docx import Document


from io import BytesIO
from docx.shared import Inches, Pt

import io
from docx.shared import Inches

def add_images_to_doc(doc, images, max_per_page=4):
    for i in range(0, len(images), max_per_page):
        chunk = images[i:i + max_per_page]

        # Calculate rows based on 2 columns
        rows = (len(chunk) + 1) // 2
        table = doc.add_table(rows=rows, cols=2)

        for index, img_obj in enumerate(chunk):
            row = index // 2
            col = index % 2

            # 1. Reset the pointer of the original image
            img_obj.seek(0)

            # 2. Read data and wrap in a new BytesIO (Solves many 'not found' issues)
            img_data = io.BytesIO(img_obj.read())

            # 3. Access cell and add picture
            cell = table.rows[row].cells[col]
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run()

            try:
                # Use a specific width to ensure it fits in the cell
                run.add_picture(img_data, width=Inches(2.2))
            except Exception as e:
                print(f"Error inserting image: {e}")

        doc.add_page_break()
    return doc
def format_indian_number(amount):
    if isinstance(amount, str):
        amount = amount.replace(",", "")   # ✅ remove commas

    amount = int(amount)
    s = str(abs(amount))

    if len(s) <= 3:
        result = s
    else:
        last3 = s[-3:]
        remaining = s[:-3]

        parts = []
        while len(remaining) > 2:
            parts.insert(0, remaining[-2:])
            remaining = remaining[:-2]

        if remaining:
            parts.insert(0, remaining)

        result = ",".join(parts) + "," + last3

    if amount < 0:
        result = "-" + result

    return result
# -------------------------------------------------
# UI CONFIG
# -------------------------------------------------
st.set_page_config(page_title="Valuation Report Generator", layout="wide")
st.title("🏠 Valuation Report Generator")

# -------------------------------------------------
# BANK + TEMPLATE CONFIG
# -------------------------------------------------
BANK_TEMPLATES = {
    "State Bank of India (SBI)": {
        "LAND_ONLY": "templates/SBI/valuation_template1.docx",
        "LAND_BUILDING": "templates/SBI/valuation_template.docx",
        "FLAT": "templates/SBI/valuation_template2.docx",
    },
       "PNB Bank": {
       "LAND_ONLY": "templates/PNB/valuation_template1.docx",
       "LAND_BUILDING": "templates/PNB/valuation_template.docx",
       "FLAT": "templates/PNB/valuation_template2.docx",
   },
    "ICICI Bank": {
        "LAND_ONLY": "templates/ICICI/valuation_template1.docx",
        "LAND_BUILDING": "templates/ICICI/valuation_template.docx",
    },
    "Federal Bank": {
        "LAND_ONLY": "templates/FEDERAL/valuation_template1.docx",
        "LAND_BUILDING": "templates/FEDERAL/valuation_template.docx",
    },
    "Union Bank of India": {
    "LAND_ONLY": "templates/UBI/valuation_template1.docx",
    "LAND_BUILDING": "templates/UBI/valuation_template.docx",
},
}

PROPERTY_TYPES = {
    "Land Only": "LAND_ONLY",
    "Land + Building": "LAND_BUILDING",
    "Flat (Composite)": "FLAT",
}

# -------------------------------------------------
# BANK & PROPERTY TYPE
# -------------------------------------------------
st.header("🏦 Bank & Property Type")

selected_bank = st.selectbox("Select Bank", list(BANK_TEMPLATES.keys()))
# Show only the property types this bank has a template for
available_types = [t for t, key in PROPERTY_TYPES.items() if key in BANK_TEMPLATES[selected_bank]]
property_type = st.selectbox("Property Type", available_types)

template_key = PROPERTY_TYPES[property_type]
template_path = os.path.join(PROJECT_ROOT, BANK_TEMPLATES[selected_bank][template_key])

if not os.path.exists(template_path):
    st.error("Template not found")
    st.stop()

# -------------------------------------------------
# FILE UPLOADS
# -------------------------------------------------
st.header("1️⃣ Upload Input Files")

uploaded_pdf = st.file_uploader("Upload Title Investigation PDF", type=["pdf"])
uploaded_docx = st.file_uploader("Upload Reference DOCX (optional)", type=["docx"])

pdf_fields, docx_fields = {}, {}

# PDF  (OCR + AI run only once per file; result kept in session_state)
if uploaded_pdf:
    pdf_key = "pdf_fields_" + uploaded_pdf.name
    if pdf_key not in st.session_state:
        os.makedirs("temp/pdf/pages", exist_ok=True)
        pdf_path = os.path.join("temp/pdf", uploaded_pdf.name)
        with open(pdf_path, "wb") as f:
            f.write(uploaded_pdf.read())

        with st.spinner("Reading PDF (OCR)..."):
            images = pdf_to_images(pdf_path, "temp/pdf/pages")
            full_text = "\n".join(ocr_english(img) for img in images)

        with st.spinner("AI extracting fields... (1-3 minutes)"):
            st.session_state[pdf_key] = extract_fields_with_ai(full_text)

    pdf_fields = st.session_state[pdf_key]

    st.subheader("📄 Extracted from PDF")
    st.json(pdf_fields)

# DOCX

if uploaded_docx:
    os.makedirs("temp/docx", exist_ok=True)
    docx_path = os.path.join("temp/docx", uploaded_docx.name)
    with open(docx_path, "wb") as f:
        f.write(uploaded_docx.read())

    docx_text = read_docx_text(docx_path)
    docx_fields = extract_all_fields(docx_text)

    st.subheader("📄 Extracted from DOCX")
    st.json(docx_fields)

# -------------------------------------------------
# MERGE FIELDS
# -------------------------------------------------
merged_fields = pdf_fields.copy()
for k, v in docx_fields.items():
    if not merged_fields.get(k) and v:
        merged_fields[k] = v

st.subheader("📑 Final OCR Fields")
st.json(merged_fields)

# -------------------------------------------------
# MANUAL OVERRIDE SECTION (KEY UPDATE)
# -------------------------------------------------
st.header("2️⃣ Verify & Edit Details")

col1, col2 = st.columns(2)

with col1:
    borrower_name = st.text_input("Borrower Name", merged_fields.get("borrower", ""))
    owner_name = st.text_input("Owner Name", merged_fields.get("owner", ""))
    survey_no = st.text_input("Survey / Resurvey No", merged_fields.get("survey_no", ""))
    village = st.text_input("Village", merged_fields.get("village", ""))
    branch = st.text_input("Branch", merged_fields.get("branch", ""))
    ref_no = st.text_input("Reference Number", merged_fields.get("ref_no", ""))
    door_no = st.text_input("Door No", merged_fields.get("door_no", ""))

with col2:
    inspection_date_dt = st.date_input("Inspection Date", datetime(2026, 1, 1))
    taluk = st.text_input("Taluk", merged_fields.get("taluk", ""))
    district = st.text_input("District", merged_fields.get("district", "Ernakulam"))
    inspection_date_str = inspection_date_dt.strftime("%d.%m.%Y")
    lat_long_input = st.text_input("Coordinates (Lat, Long)", "")
    uploaded_files = st.file_uploader("Upload Property Photos", type=["jpg", "jpeg", "png"],accept_multiple_files=True )
    if uploaded_files:

        st.write(f"{len(uploaded_files)} images uploaded")
st.subheader("🧭 Boundaries")

b1, b2 = st.columns(2)
doa_date = (inspection_date_dt - timedelta(days=1)).strftime("%d.%m.%Y")
with b1:
    boundary_east = st.text_input("East", merged_fields.get("boundary_east", ""))
    boundary_west = st.text_input("West", merged_fields.get("boundary_west", ""))

with b2:
    boundary_north = st.text_input("North", merged_fields.get("boundary_north", ""))
    boundary_south = st.text_input("South", merged_fields.get("boundary_south", ""))

st.subheader("📜 Legal Documents")

d1, d2 = st.columns(2)
with d1:
    sale_deed = st.text_input("Sale Deed No.", merged_fields.get("sale_deed", ""))
    location_sketch = st.text_input("Location Sketch No.", merged_fields.get("location_sketch", ""))
with d2:
    tax_receipt = st.text_input("Land Tax Receipt No.", merged_fields.get("tax_receipt", ""))
    possession_cert = st.text_input("Possession Certificate No.", merged_fields.get("possession_cert", ""))

# -------------------------------------------------
# FLAT DETAILS (only for Flat (Composite))
# -------------------------------------------------
flat = {}
if property_type == "Flat (Composite)":
    st.subheader("🏢 Flat Details")
    f1, f2 = st.columns(2)
    with f1:
        flat["UDS_EXTENT"] = st.text_input("Undivided share (e.g. 0.77386 Ares)")
        flat["TOTAL_LAND_AREA"] = st.text_input("Total land area of project")
        flat["PARENT_DEEDS"] = st.text_input("Parent deeds (e.g. Sale Deed Nos. 4787/2007 ... of Ernakulam SRO)")
        flat["APARTMENT_NO"] = st.text_input("Apartment No.")
        flat["FLOOR"] = st.text_input("Floor (e.g. 19th)")
        flat["BLOCK"] = st.text_input("Block / Wing")
        flat["PROJECT_NAME"] = st.text_input("Project name")
    with f2:
        flat["SUPER_BUILTUP_AREA"] = st.text_input("Super built-up area (Sq. m)")
        flat["CARPET_AREA"] = st.text_input("Carpet area (Sq. m)")
        flat["BALCONY_AREA"] = st.text_input("Balcony / verandah area (Sq. m)")
        flat["COMMON_AREA"] = st.text_input("Share of common area (Sq. m)")
        flat["CAR_PARKING"] = st.text_input("Car parking", "one covered car parking space")
        flat["PRESENT_OWNER"] = st.text_input("Present owner (seller)")

# -------------------------------------------------
# VALUATION INPUTS
# -------------------------------------------------
st.header("3️⃣ Valuation Inputs")

extent_are = st.number_input("Extent of land (Ares)", 0.0, step=0.01)
land_rate = st.number_input("Land rate per cent (₹)", 0)

if property_type != "Land Only":
    building_sqft = st.number_input("Building area (Sqft)", 0.0)
    building_rate = st.number_input("Building rate per Sqft (₹)", 0)
    building_age = st.number_input("Building age (Years)", 0)
else:
    building_sqft = building_rate = building_age = 0
current_year = datetime.now().year
year_of_construction = current_year - building_age
# -------------------------------------------------
# GENERATE REPORT
# -------------------------------------------------
if st.button("📄 Generate Valuation Report"):

    total_cents = round(extent_are * ARE_TO_CENT, 2)
    total_land_value = round_down(total_cents * land_rate, 1000)
    dep = 0
    gross_building_value = 0
    depreciated_building_value = 0
    yoc_display = "N/A"
    if property_type != "Land Only":
        gross_building_value = building_sqft * building_rate
        dep = building_age * DEPRECIATION_PER_YEAR
        depreciated_building_value = round_down(gross_building_value * (1 - dep), 1000)
    else:
        gross_building_value = depreciated_building_value = 0

    federal_vals = {}
    if selected_bank  in ["Federal Bank", "PNB"] and property_type == "Land + Building":
        federal_vals = calculate_federal_separate(total_land_value, depreciated_building_value)
        market_value = federal_vals["MARKET_VALUE"]
        realizable_value = federal_vals["REALIZABLE_VALUE"]
        distress_value = federal_vals["DISTRESS_VALUE"]
    else:
        market_value = round_down(total_land_value + depreciated_building_value, 100000)
        realizable_value = round_down(market_value * 0.90, 100000)
        distress_value = round_down(market_value * 0.80, 100000)

    context = {
        "REF_NO": ref_no,
        "BORROWER_NAME": borrower_name,
        "OWNER_NAME":owner_name,
        "SURVEY_NO": survey_no,
        "VILLAGE": village,
        "TALUK": taluk,
        "DISTRICT": district,
        "DOOR_NO": door_no,
        "INSPECTION_DATE": inspection_date_str,
        "YOC": year_of_construction,
        "DOA_DATE": doa_date,
        "LAT_LONG": lat_long_input,
        "BOUNDARY_EAST": boundary_east,
        "BOUNDARY_WEST": boundary_west,
        "BOUNDARY_NORTH": boundary_north,
        "BOUNDARY_SOUTH": boundary_south,
        "SALE_DEED": sale_deed,
        "LOCATION_SKETCH": location_sketch,
        "TAX_RECEIPT": tax_receipt,
        "POSSESSION_CERT": possession_cert,
        "Branch": branch,
        "EXTENT": f"{extent_are} Ares",
        "TOTAL_CENTS": total_cents,
        "LAND_RATE_PER_CENT": land_rate,
        "RAW_LAND_VALUE": format_indian_number(total_cents * land_rate),
        "TOTAL_LAND_VALUE": total_land_value,
        "BAS": building_sqft,
        "BRPS": format_indian_number(building_rate),
        "BAY": building_age,
        "BGV": format_indian_number(gross_building_value),
        "DBV": format_indian_number(depreciated_building_value),
        "DPN": format_indian_number(gross_building_value * dep),
        "MARKET_VALUE": format_indian_number(market_value),
        "REALIZABLE_VALUE": format_indian_number(realizable_value),
        "DISTRESS_VALUE": format_indian_number(distress_value),
        "MARKET_VALUE_WORDS": to_words(market_value),
        "REALIZABLE_VALUE_WORDS": to_words(realizable_value),
        "DISTRESS_VALUE_WORDS": to_words(distress_value),
        **flat,
        **federal_vals,
        "PLACE": district,
        "REPORT_DATE": datetime.now().strftime("%d-%m-%Y")
       }

    # 1. Generate the report from template first
    os.makedirs("output", exist_ok=True)
    output_path = "output/valuation_report.docx"

    # This usually returns a document or saves it to output_path
    generate_report(template_path, output_path, context)

    # 2. Now open THAT generated document and append photos
    if uploaded_files:
        # Load the document that was just created by generate_report
        doc = Document(output_path)

        doc.add_page_break()
        title = doc.add_paragraph()
        run = title.add_run("PROPERTY PHOTOS")
        run.bold = True
        run.font.size = Pt(14)

        # Use your function to add images to the report document
        doc = add_images_to_doc(doc, uploaded_files, max_per_page=4)

        # 3. SAVE the document again to capture the images
        doc.save(output_path)

    st.success("✅ Valuation Report Generated with Photos")
# ...
    with open(output_path, "rb") as f:
        st.download_button("⬇️ Download Report", f, "valuation_report.docx")